"""
Render synthetic profiles into real documents (PDF / DOCX / TXT / MD / HTML)
with many layouts, so the SLM is trained on exactly what text_extractor.py
produces from real-world files.

Templates
    pdf_classic   single column, varied entry patterns, right-aligned dates
    pdf_sidebar   two columns (sidebar left or right, optional background)
    pdf_form      Vietnamese-HR style ruled form (grid tables)
    docx_classic  Heading styles / bold labels, list bullets, tab-aligned dates
    docx_layout   template-style 2-column layout table
    docx_form     application form built from grid tables
    txt / md / html
"""
from __future__ import annotations

import os
import random
from glob import glob
from pathlib import Path
from xml.sax.saxutils import escape

from training.synth import Profile

# ---------------------------------------------------------------------------
# Fonts (must cover Vietnamese; reportlab's built-in Helvetica does not)
# ---------------------------------------------------------------------------

FONT_CANDIDATES = [  # (family, regular, bold, italic)
    ("DejaVuSans", "DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf"),
    ("DejaVuSerif", "DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf"),
    ("LiberationSans", "LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf"),
    ("LiberationSerif", "LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf"),
    ("FreeSans", "FreeSans.ttf", "FreeSansBold.ttf", "FreeSansOblique.ttf"),
    ("FreeSerif", "FreeSerif.ttf", "FreeSerifBold.ttf", "FreeSerifItalic.ttf"),
    ("NotoSans", "NotoSans-Regular.ttf", "NotoSans-Bold.ttf", "NotoSans-Italic.ttf"),
    ("Arial", "arial.ttf", "arialbd.ttf", "ariali.ttf"),
]
FONT_DIRS = [os.environ.get("CV_FONTS_DIR", ""), "/usr/share/fonts", "/usr/local/share/fonts",
             os.path.expanduser("~/.fonts"), "/Library/Fonts", "C:/Windows/Fonts"]
DOCX_FONTS = ["Arial", "Calibri", "Times New Roman", "Cambria", "Segoe UI", "Tahoma", "Roboto"]

_registered: list[str] | None = None


def _find_font(filename: str) -> str | None:
    for d in FONT_DIRS:
        if d and os.path.isdir(d):
            hits = glob(os.path.join(d, "**", filename), recursive=True)
            if hits:
                return hits[0]
    try:  # matplotlib ships DejaVu
        import matplotlib
        p = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / filename
        if p.exists():
            return str(p)
    except ImportError:
        pass
    return None


def pdf_fonts() -> list[str]:
    """Register every available Unicode TTF family with reportlab."""
    global _registered
    if _registered is not None:
        return _registered
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    _registered = []
    for family, reg, bold, ital in FONT_CANDIDATES:
        reg_p, bold_p = _find_font(reg), _find_font(bold)
        if not reg_p or not bold_p:
            continue
        ital_p = _find_font(ital) or reg_p
        pdfmetrics.registerFont(TTFont(family, reg_p))
        pdfmetrics.registerFont(TTFont(f"{family}-Bold", bold_p))
        pdfmetrics.registerFont(TTFont(f"{family}-Italic", ital_p))
        pdfmetrics.registerFont(TTFont(f"{family}-BoldItalic", bold_p))
        addMapping(family, 0, 0, family)
        addMapping(family, 1, 0, f"{family}-Bold")
        addMapping(family, 0, 1, f"{family}-Italic")
        addMapping(family, 1, 1, f"{family}-BoldItalic")
        _registered.append(family)
    if not _registered:
        raise RuntimeError(
            "No Unicode TTF font found for PDF rendering (Vietnamese needs one). "
            "Install fonts-dejavu / fonts-liberation, or set CV_FONTS_DIR to a folder with DejaVuSans.ttf."
        )
    return _registered


# ---------------------------------------------------------------------------
# Logical content (shared by all templates)
# ---------------------------------------------------------------------------

HEADER_FIELDS = ("email", "phone", "address", "city", "country", "linkedin", "website")
PERSONAL_BLOCK_FIELDS = ("date_of_birth", "gender", "place_of_birth", "nationality", "marital_status",
                         "id_number", "id_issue_date", "id_issue_place")


def contact_items(p: Profile, rng: random.Random, icons: bool) -> list[tuple[str, str, str]]:
    """(key, label, value) for header contact fields."""
    icon = {"email": "✉ ", "phone": "☎ ", "address": "⌂ ", "city": "⌂ "} if icons else {}
    out = []
    for k in HEADER_FIELDS:
        if k in p.personal_disp:
            out.append((k, p.label(k), icon.get(k, "") + p.personal_disp[k]))
    return out


def personal_items(p: Profile, include_contact: bool) -> list[tuple[str, str, str]]:
    keys = (["full_name"] if include_contact else []) + list(PERSONAL_BLOCK_FIELDS) + \
        (list(HEADER_FIELDS) if include_contact else [])
    return [(k, p.label(k), p.personal_disp[k]) for k in keys if k in p.personal_disp]


def section_order(p: Profile, rng: random.Random) -> list[str]:
    main = ["experience", "education"] if rng.random() < 0.6 else ["education", "experience"]
    rest = ["languages", "certificates", "awards", "courses", "family"]
    rng.shuffle(rest)
    order = []
    if "summary" in p.extras and rng.random() < 0.85:
        order.append("summary")
    order += main
    if "skills" in p.extras:
        rest.insert(rng.randint(0, len(rest)), "skills")
    if "projects" in p.extras:
        rest.insert(rng.randint(0, len(rest)), "projects")
    order += rest
    for extra in ("interests", "references"):
        if extra in p.extras:
            order.append(extra)
    if "summary" in p.extras and "summary" not in order:
        order.append("summary")
    return [s for s in order if p.sections.get(s) or s in p.extras]


def entry_lines(sec: str, e, p: Profile, pattern: int) -> list[tuple]:
    """Linear rendering of one entry as typed lines:
    ("title", t) ("text", t) ("meta", t) ("bullet", t) ("lr", left, right)"""
    d = e.disp
    L = p.flabel
    if sec == "experience":
        loc = f", {d['location']}" if d.get("location") else ""
        if pattern == 0:
            lines = [("title", f"{d['job_title']} | {d['company']}{loc}"), ("meta", d["dates"])]
        elif pattern == 1:
            lines = [("lr", f"<b>{escape(d['company'])}</b>{escape(loc)}", escape(d["dates"])), ("text", d["job_title"])]
        elif pattern == 2:
            lines = [("title", f"{d['dates']}: {d['job_title']} - {d['company']}{loc}")]
        elif pattern == 3:
            lines = [("lr", f"<b>{escape(d['job_title'])}</b>", escape(d["dates"])), ("text", f"{d['company']}{loc}")]
        else:
            lines = [("kv", L("company"), d["company"]), ("kv", L("job_title"), d["job_title"]),
                     ("kv", L("time"), d["dates"])]
            if d.get("location"):
                lines.append(("kv", L("location"), d["location"]))
            if d["bullets"]:
                lines.append(("text", f"{L('description')}:"))
        lines += [("bullet", b) for b in d["bullets"]]
        return lines
    if sec == "education":
        dates = f"{d['start']} - {d['end']}" if d.get("start") else d["end"]
        deg = " ".join(x for x in (d.get("degree"), d.get("major")) if x)
        grade = f"{L('grade')}: {d['grade']}" if d.get("grade") else ""
        if pattern in (0, 3):
            lines = [("lr", f"<b>{escape(d['institution'])}</b>", escape(dates)), ("text", deg)]
        elif pattern == 1:
            lines = [("title", deg or d["institution"]), ("text", f"{d['institution']} | {dates}")] if deg \
                else [("title", d["institution"]), ("text", dates)]
        elif pattern == 2:
            lines = [("title", f"{dates}: {d['institution']}"),
                     ("text", f"{L('major')}: {d['major']}" + (f" | {L('degree')}: {d['degree']}" if d.get("degree") else ""))]
        else:
            lines = [("kv", L("institution"), d["institution"]), ("kv", L("major"), d["major"])]
            if d.get("degree"):
                lines.append(("kv", L("degree"), d["degree"]))
            lines.append(("kv", L("time"), dates))
        if grade:
            lines.append(("text", grade))
        return lines
    if sec == "languages":
        return [("bullet", f"{d['language']}: {d['proficiency']}" if d.get("proficiency") and pattern % 2
                 else f"{d['language']} ({d['proficiency']})" if d.get("proficiency") else d["language"])]
    if sec in ("certificates", "awards", "courses"):
        issuer = d.get("issuer") or d.get("provider")
        parts = [d["name"]]
        if issuer:
            parts.append(issuer)
        text = " - ".join(parts) if pattern % 2 else ", ".join(parts)
        if d.get("date"):
            text = f"{d['date']}: {text}" if pattern == 2 else f"{text} ({d['date']})"
        return [("bullet", text)]
    if sec == "family":
        bits = [d["full_name"], d["relationship"]]
        if d.get("yob"):
            bits.append(f"{p.flabel('yob')}: {d['yob']}")
        if d.get("occupation"):
            bits.append(d["occupation"])
        if d.get("phone"):
            bits.append(d["phone"])
        return [("bullet", " - ".join(bits))]
    return []


def extra_lines(sec: str, p: Profile) -> list[tuple]:
    v = p.extras[sec]
    if sec == "summary":
        return [("text", v)]
    if sec == "skills":
        return [("text", ", ".join(v))] if len(v) > 5 else [("bullet", s) for s in v]
    return [("bullet", s) for s in v]


def section_lines(sec: str, p: Profile, rng: random.Random, pattern: int) -> list[tuple]:
    if sec in p.extras and sec not in p.sections:
        return extra_lines(sec, p)
    out = []
    for e in p.sections.get(sec, []):
        out += entry_lines(sec, e, p, pattern)
        out.append(("space",))
    return out


def _plain(text: str) -> str:
    return text.replace("<b>", "").replace("</b>", "").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def _pdf_styles(rng: random.Random):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    font = rng.choice(pdf_fonts())
    body = rng.choice([9, 9.5, 10, 10.5, 11])
    accent = HexColor(rng.choice(["#000000", "#1a3c6e", "#2c3e50", "#7a1f1f", "#0b5e4a", "#333333"]))
    st = {
        "name": ParagraphStyle("name", fontName=f"{font}-Bold", fontSize=rng.choice([18, 20, 22, 26]),
                               leading=30, textColor=accent, spaceAfter=4),
        "title": ParagraphStyle("title", fontName=font, fontSize=body + 2, leading=body + 5, spaceAfter=4),
        "h": ParagraphStyle("h", fontName=f"{font}-Bold", fontSize=body + rng.choice([1, 2, 3]),
                            leading=body + 7, textColor=accent, spaceBefore=rng.choice([6, 9, 12]), spaceAfter=4),
        "entry": ParagraphStyle("entry", fontName=f"{font}-Bold", fontSize=body + 0.5, leading=body + 4, spaceBefore=2),
        "text": ParagraphStyle("text", fontName=font, fontSize=body, leading=body + 3.5),
        "meta": ParagraphStyle("meta", fontName=f"{font}-Italic", fontSize=body - 0.5, leading=body + 3,
                               textColor=HexColor("#555555")),
        "bullet": ParagraphStyle("bullet", fontName=font, fontSize=body, leading=body + 3.5, leftIndent=12,
                                 bulletIndent=2, bulletFontName=font),
        "right": ParagraphStyle("right", fontName=font, fontSize=body, leading=body + 3.5, alignment=TA_RIGHT),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=body - 0.5, leading=body + 2.5),
        "cellb": ParagraphStyle("cellb", fontName=f"{font}-Bold", fontSize=body - 0.5, leading=body + 2.5),
    }
    return st, accent, font


def _pdf_flow(lines: list[tuple], st: dict, width: float, rng: random.Random) -> list:
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    bullet_char = rng.choice(["•", "•", "-", "▪", "●", "–", "✓"])
    out = []
    for ln in lines:
        kind = ln[0]
        if kind == "space":
            out.append(Spacer(1, 3))
        elif kind == "title":
            out.append(Paragraph(escape(ln[1]), st["entry"]))
        elif kind == "text":
            out.append(Paragraph(escape(ln[1]), st["text"]))
        elif kind == "meta":
            out.append(Paragraph(escape(ln[1]), st["meta"]))
        elif kind == "bullet":
            out.append(Paragraph(escape(ln[1]), st["bullet"], bulletText=bullet_char))
        elif kind == "kv":
            out.append(Paragraph(f"<b>{escape(ln[1])}:</b> {escape(ln[2])}", st["text"]))
        elif kind == "lr":
            t = Table([[Paragraph(ln[1], st["text"]), Paragraph(ln[2], st["right"])]],
                      colWidths=[width * 0.66, width * 0.34])
            t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            out.append(t)
    return out


def _heading(p: Profile, sec: str, st: dict, accent, rng: random.Random, rule: bool) -> list:
    from reportlab.platypus import HRFlowable, Paragraph
    out = [Paragraph(escape(p.heading[sec]), st["h"])]
    if rule:
        out.append(HRFlowable(width="100%", thickness=0.8, color=accent, spaceBefore=0, spaceAfter=4))
    return out


def _header_flow(p: Profile, st: dict, rng: random.Random, contacts_inline: bool) -> list:
    from reportlab.platypus import Paragraph
    out = [Paragraph(escape(p.personal_disp["full_name"]), st["name"])]
    if p.personal_disp.get("current_title"):
        out.append(Paragraph(escape(p.personal_disp["current_title"]), st["title"]))
    items = contact_items(p, rng, icons=rng.random() < 0.2)
    labelled = rng.random() < 0.4
    vals = [f"{lab}: {v}" if labelled else v for _, lab, v in items]
    if contacts_inline and vals:
        sep = rng.choice([" | ", "  •  ", " · ", "   "])
        out.append(Paragraph(escape(sep.join(vals)), st["text"]))
    else:
        out += [Paragraph(escape(v), st["text"]) for v in vals]
    return out


def render_pdf_classic(p: Profile, path: str, rng: random.Random):
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer

    size = rng.choice([A4, A4, LETTER])
    margin = rng.choice([1.5, 2.0, 2.5]) * cm
    doc = SimpleDocTemplate(path, pagesize=size, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin)
    width = size[0] - 2 * margin
    st, accent, _ = _pdf_styles(rng)
    rule = rng.random() < 0.6
    pattern = rng.randint(0, 4)
    story = _header_flow(p, st, rng, contacts_inline=rng.random() < 0.5)
    story.append(Spacer(1, 6))
    personal = personal_items(p, include_contact=False)
    if personal:
        story += _heading(p, "personal", st, accent, rng, rule)
        story += _pdf_flow([("kv", lab, v) for _, lab, v in personal], st, width, rng)
    for sec in section_order(p, rng):
        story += _heading(p, sec, st, accent, rng, rule)
        story += _pdf_flow(section_lines(sec, p, rng, pattern), st, width, rng)
    _add_page_numbers(doc, story, rng)


def _add_page_numbers(doc, story, rng: random.Random, **kw):
    style = rng.choice([None, "Page {n}", "Trang {n}", "{n}", "Page {n} of 2"])
    font = pdf_fonts()[0]

    def on_page(canvas, d):
        if style:
            canvas.setFont(font, 8)
            canvas.drawRightString(d.pagesize[0] - 40, 25, style.format(n=d.page))
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page, **kw)


def render_pdf_sidebar(p: Profile, path: str, rng: random.Random):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import BaseDocTemplate, Frame, FrameBreak, PageTemplate, Spacer

    W, H = A4
    m = 28
    side_w = rng.choice([0.30, 0.33, 0.36]) * (W - 2 * m)
    gap = rng.choice([10, 16, 24])
    main_w = W - 2 * m - side_w - gap
    left_sidebar = rng.random() < 0.75
    sx = m if left_sidebar else m + main_w + gap
    mx = m + side_w + gap if left_sidebar else m
    frames = [Frame(sx, m, side_w, H - 2 * m, id="side", leftPadding=4, rightPadding=4),
              Frame(mx, m, main_w, H - 2 * m, id="main", leftPadding=4, rightPadding=4)]
    bg = rng.choice([None, "#f0f3f7", "#e8eef5", "#f5efe6"])
    font = pdf_fonts()[0]

    def draw_bg(canvas, d):
        if bg:
            canvas.setFillColor(HexColor(bg))
            canvas.rect(sx - 6, 0, side_w + 12, H, stroke=0, fill=1)
        canvas.setFont(font, 7)

    doc = BaseDocTemplate(path, pagesize=A4, leftMargin=m, rightMargin=m, topMargin=m, bottomMargin=m)
    later = [Frame(m, m, W - 2 * m, H - 2 * m, id="full")]
    doc.addPageTemplates([PageTemplate(id="first", frames=frames, onPage=draw_bg, autoNextPageTemplate="later"),
                          PageTemplate(id="later", frames=later)])
    st, accent, _ = _pdf_styles(rng)
    rule = rng.random() < 0.5
    pattern = rng.randint(0, 4)
    order = section_order(p, rng)
    side_secs = [s for s in order if s in ("languages", "skills", "certificates", "interests", "awards", "references")
                 and rng.random() < 0.85]
    main_secs = [s for s in order if s not in side_secs]

    side, main = [], []
    name_in_side = rng.random() < 0.5
    header = _header_flow(p, st, rng, contacts_inline=False)
    if name_in_side:
        side += header
    else:
        main += header[:2 if p.personal_disp.get("current_title") else 1]
        side += [f for f in header[2 if p.personal_disp.get("current_title") else 1:]]
    side.append(Spacer(1, 8))
    personal = personal_items(p, include_contact=False)
    if personal:
        side += _heading(p, "personal", st, accent, rng, rule)
        side += _pdf_flow([("text", f"{lab}: {v}") for _, lab, v in personal], st, side_w - 8, rng)
    for sec in side_secs:
        side += _heading(p, sec, st, accent, rng, rule)
        side += _pdf_flow(section_lines(sec, p, rng, pattern), st, side_w - 8, rng)
    for sec in main_secs:
        main += _heading(p, sec, st, accent, rng, rule)
        main += _pdf_flow(section_lines(sec, p, rng, pattern), st, main_w - 8, rng)
    first, second = (side, main) if frames[0].id == "side" else (main, side)
    doc.build(first + [FrameBreak()] + second)


def render_pdf_form(p: Profile, path: str, rng: random.Random):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    W = A4[0] - 3 * cm
    st, accent, _ = _pdf_styles(rng)
    vi = p.lang in ("vi", "bi")
    title = rng.choice(["PHIẾU THÔNG TIN ỨNG VIÊN", "SƠ YẾU LÝ LỊCH", "ĐƠN ỨNG TUYỂN"] if vi
                       else ["APPLICATION FORM", "CANDIDATE INFORMATION FORM"])
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, HexColor("#444444")),
                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                       ("BACKGROUND", (0, 0), (-1, 0), HexColor("#dde3ea"))])
    kv_grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, HexColor("#444444")),
                          ("VALIGN", (0, 0), (-1, -1), "TOP")])
    C = lambda t, b=False: Paragraph(escape(t), st["cellb" if b else "cell"])  # noqa: E731
    story = [Paragraph(escape(title), st["name"]), Spacer(1, 6)]
    if p.personal_disp.get("current_title"):
        story.append(Paragraph(escape(f"{p.label('current_title')}: {p.personal_disp['current_title']}"), st["text"]))
    items = personal_items(p, include_contact=True)
    story += _heading(p, "personal", st, accent, rng, False)
    rows = []
    two_per_row = rng.random() < 0.6
    step = 2 if two_per_row else 1
    for i in range(0, len(items), step):
        row = []
        for _, lab, v in items[i:i + step]:
            row += [C(lab, True), C(v)]
        while len(row) < 2 * step:
            row += [C(""), C("")]
        rows.append(row)
    widths = [W * 0.17, W * 0.33, W * 0.17, W * 0.33] if two_per_row else [W * 0.3, W * 0.7]
    t = Table(rows, colWidths=widths)
    t.setStyle(kv_grid)
    story += [t, Spacer(1, 8)]
    L = p.flabel
    for sec in section_order(p, rng):
        story += _heading(p, sec, st, accent, rng, False)
        entries = p.sections.get(sec) or []
        if sec == "experience" and entries:
            rows = [[C(L("from"), True), C(L("to"), True), C(L("company"), True), C(L("job_title"), True),
                     C(L("description"), True)]]
            for e in entries:
                d = e.disp
                desc = Paragraph("<br/>".join("- " + escape(b) for b in d["bullets"]), st["cell"])
                rows.append([C(d["start"]), C(d["end"]), C(d["company"] + (f", {d['location']}" if d.get("location") else "")),
                             C(d["job_title"]), desc])
            widths = [W * 0.14, W * 0.14, W * 0.22, W * 0.17, W * 0.33]
        elif sec == "education" and entries:
            rows = [[C(L("from"), True), C(L("to"), True), C(L("institution"), True), C(L("major"), True),
                     C(L("degree"), True), C(L("grade"), True)]]
            for e in entries:
                d = e.disp
                rows.append([C(d.get("start", "")), C(d["end"]), C(d["institution"]), C(d["major"]),
                             C(d.get("degree", "")), C(d.get("grade", ""))])
            widths = [W * 0.13, W * 0.13, W * 0.28, W * 0.18, W * 0.14, W * 0.14]
        elif sec == "family" and entries:
            rows = [[C(L("name"), True), C(L("relationship"), True), C(L("yob"), True), C(L("occupation"), True)]]
            for e in entries:
                d = e.disp
                job = " - ".join(x for x in (d.get("occupation"), d.get("phone")) if x)
                rows.append([C(d["full_name"]), C(d["relationship"]), C(d.get("yob", "")), C(job)])
            widths = [W * 0.35, W * 0.2, W * 0.15, W * 0.3]
        elif sec == "languages" and entries:
            rows = [[C(L("language"), True), C(L("proficiency"), True)]]
            rows += [[C(e.disp["language"]), C(e.disp.get("proficiency", ""))] for e in entries]
            widths = [W * 0.4, W * 0.6]
        elif sec in ("certificates", "awards", "courses") and entries:
            rows = [[C(L("name"), True), C(L("issuer"), True), C(L("date"), True)]]
            rows += [[C(e.disp["name"]), C(e.disp.get("issuer") or e.disp.get("provider") or ""),
                      C(e.disp.get("date", ""))] for e in entries]
            widths = [W * 0.5, W * 0.3, W * 0.2]
        else:
            story += _pdf_flow(section_lines(sec, p, rng, 0), st, W, rng)
            continue
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(grid)
        story += [t, Spacer(1, 6)]
    if rng.random() < 0.5:
        story += [Spacer(1, 14), Paragraph(escape("Tôi cam đoan những thông tin trên là đúng sự thật." if vi else
                                                 "I certify that the information above is true."), st["meta"])]
    _add_page_numbers(doc, story, rng)


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def _docx_base(rng: random.Random):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = rng.choice(DOCX_FONTS)
    style.font.size = Pt(rng.choice([10, 10.5, 11, 12]))
    return d


def _docx_heading(d, text: str, rng: random.Random, mode: int, container=None):
    from docx.shared import Pt
    target = container or d
    if mode == 0 and container is None:
        d.add_heading(text, level=rng.choice([1, 2]))
        return
    para = target.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(rng.choice([12, 13, 14]))


def _docx_lines(target, lines: list[tuple], rng: random.Random, doc=None):
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm
    use_list_style = doc is not None and rng.random() < 0.6
    bullet = rng.choice(["• ", "- ", "+ ", "* "])
    for ln in lines:
        kind = ln[0]
        if kind == "space":
            continue
        if kind == "bullet":
            if use_list_style:
                target.add_paragraph(ln[1], style="List Bullet")
            else:
                target.add_paragraph(bullet + ln[1])
        elif kind == "title":
            target.add_paragraph().add_run(ln[1]).bold = True
        elif kind == "meta":
            target.add_paragraph().add_run(ln[1]).italic = True
        elif kind == "kv":
            para = target.add_paragraph()
            para.add_run(f"{ln[1]}: ").bold = True
            para.add_run(ln[2])
        elif kind == "lr":
            para = target.add_paragraph()
            if doc is not None:
                para.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
            left = _plain(ln[1])
            para.add_run(left).bold = "<b>" in ln[1]
            para.add_run("\t" + _plain(ln[2]))
        else:
            target.add_paragraph(ln[1])


def render_docx_classic(p: Profile, path: str, rng: random.Random):
    from docx.shared import Pt
    d = _docx_base(rng)
    name_para = d.add_paragraph()
    run = name_para.add_run(p.personal_disp["full_name"])
    run.bold = True
    run.font.size = Pt(rng.choice([18, 20, 24]))
    if p.personal_disp.get("current_title"):
        d.add_paragraph(p.personal_disp["current_title"])
    items = contact_items(p, rng, icons=False)
    if rng.random() < 0.3:  # contact block in the page header (common in Word templates)
        d.sections[0].header.paragraphs[0].text = " | ".join(v for _, _, v in items)
    elif rng.random() < 0.5:
        d.add_paragraph(" | ".join(v for _, _, v in items))
    else:
        for _, lab, v in items:
            d.add_paragraph(f"{lab}: {v}")
    mode = rng.randint(0, 1)
    pattern = rng.randint(0, 4)
    personal = personal_items(p, include_contact=False)
    if personal:
        _docx_heading(d, p.heading["personal"], rng, mode)
        _docx_lines(d, [("kv", lab, v) for _, lab, v in personal], rng, doc=d)
    for sec in section_order(p, rng):
        _docx_heading(d, p.heading[sec], rng, mode)
        _docx_lines(d, section_lines(sec, p, rng, pattern), rng, doc=d)
    d.save(path)


def render_docx_layout(p: Profile, path: str, rng: random.Random):
    from docx.shared import Pt
    d = _docx_base(rng)
    table = d.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    side, main = (left, right) if rng.random() < 0.8 else (right, left)
    r = side.paragraphs[0].add_run(p.personal_disp["full_name"])
    r.bold, r.font.size = True, Pt(18)
    if p.personal_disp.get("current_title"):
        side.add_paragraph(p.personal_disp["current_title"])
    for _, lab, v in contact_items(p, rng, icons=False):
        side.add_paragraph(v)
    pattern = rng.randint(0, 4)
    order = section_order(p, rng)
    side_secs = [s for s in order if s in ("languages", "skills", "certificates", "interests")]
    personal = personal_items(p, include_contact=False)
    if personal:
        _docx_heading(d, p.heading["personal"], rng, 1, container=side)
        _docx_lines(side, [("text", f"{lab}: {v}") for _, lab, v in personal], rng)
    for sec in side_secs:
        _docx_heading(d, p.heading[sec], rng, 1, container=side)
        _docx_lines(side, section_lines(sec, p, rng, pattern), rng)
    first = True
    for sec in order:
        if sec in side_secs:
            continue
        if first:
            main.paragraphs[0].add_run(p.heading[sec]).bold = True
            first = False
        else:
            _docx_heading(d, p.heading[sec], rng, 1, container=main)
        _docx_lines(main, section_lines(sec, p, rng, pattern), rng)
    d.save(path)


def render_docx_form(p: Profile, path: str, rng: random.Random):
    d = _docx_base(rng)
    vi = p.lang in ("vi", "bi")
    d.add_heading(rng.choice(["PHIẾU THÔNG TIN ỨNG VIÊN", "SƠ YẾU LÝ LỊCH"] if vi else ["APPLICATION FORM"]), 0)
    L = p.flabel

    def grid(rows):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                t.rows[i].cells[j].text = val
        d.add_paragraph("")

    if p.personal_disp.get("current_title"):
        d.add_paragraph(f"{p.label('current_title')}: {p.personal_disp['current_title']}")
    d.add_paragraph().add_run(p.heading["personal"]).bold = True
    items = personal_items(p, include_contact=True)
    grid([[lab, v] for _, lab, v in items])
    for sec in section_order(p, rng):
        entries = p.sections.get(sec) or []
        d.add_paragraph().add_run(p.heading[sec]).bold = True
        if sec == "experience" and entries:
            rows = [[L("from"), L("to"), L("company"), L("job_title"), L("description")]]
            rows += [[e.disp["start"], e.disp["end"],
                      e.disp["company"] + (f", {e.disp['location']}" if e.disp.get("location") else ""),
                      e.disp["job_title"], "\n".join(e.disp["bullets"])] for e in entries]
            grid(rows)
        elif sec == "education" and entries:
            rows = [[L("from"), L("to"), L("institution"), L("major"), L("degree")]]
            rows += [[e.disp.get("start", ""), e.disp["end"], e.disp["institution"], e.disp["major"],
                      e.disp.get("degree", "")] for e in entries]
            grid(rows)
            for e in entries:
                if e.disp.get("grade"):
                    d.add_paragraph(f"{e.disp['institution']} - {L('grade')}: {e.disp['grade']}")
        elif sec == "family" and entries:
            rows = [[L("name"), L("relationship"), L("yob"), L("occupation")]]
            rows += [[e.disp["full_name"], e.disp["relationship"], e.disp.get("yob", ""),
                      " - ".join(x for x in (e.disp.get("occupation"), e.disp.get("phone")) if x)]
                     for e in entries]
            grid(rows)
        else:
            _docx_lines(d, section_lines(sec, p, rng, rng.randint(0, 3)), rng, doc=d)
    d.save(path)


# ---------------------------------------------------------------------------
# Plain formats
# ---------------------------------------------------------------------------

def _text_lines(p: Profile, rng: random.Random, markdown: bool) -> list[str]:
    out = [("# " if markdown else "") + p.personal_disp["full_name"]]
    if p.personal_disp.get("current_title"):
        out.append(p.personal_disp["current_title"])
    out += [f"{lab}: {v}" if rng.random() < 0.5 else v for _, lab, v in contact_items(p, rng, False)]
    pattern = rng.randint(0, 4)
    bullet = rng.choice(["- ", "* ", "• "])
    blocks = []
    personal = personal_items(p, include_contact=False)
    if personal:
        blocks.append(("personal", [("kv", lab, v) for _, lab, v in personal]))
    for sec in section_order(p, rng):
        blocks.append((sec, section_lines(sec, p, rng, pattern)))
    for sec, lines in blocks:
        out.append("")
        out.append(("## " if markdown else "") + (p.heading[sec].upper() if not markdown else p.heading[sec]))
        for ln in lines:
            if ln[0] == "space":
                out.append("")
            elif ln[0] == "bullet":
                out.append(bullet + ln[1])
            elif ln[0] == "kv":
                out.append(f"{ln[1]}: {ln[2]}")
            elif ln[0] == "lr":
                out.append(f"{_plain(ln[1])}{rng.choice(['    ', ' | ', ' — ', chr(9)])}{_plain(ln[2])}")
            else:
                out.append(ln[1])
    return out


def render_txt(p: Profile, path: str, rng: random.Random):
    Path(path).write_text("\n".join(_text_lines(p, rng, markdown=False)), encoding="utf-8")


def render_md(p: Profile, path: str, rng: random.Random):
    Path(path).write_text("\n".join(_text_lines(p, rng, markdown=True)), encoding="utf-8")


def render_html(p: Profile, path: str, rng: random.Random):
    h = ["<html><head><meta charset='utf-8'><title>CV</title></head><body>",
         f"<h1>{escape(p.personal_disp['full_name'])}</h1>"]
    if p.personal_disp.get("current_title"):
        h.append(f"<p>{escape(p.personal_disp['current_title'])}</p>")
    for k, lab, v in contact_items(p, rng, False):
        if k == "email":
            h.append(f"<p><a href='mailto:{escape(v)}'>{escape(v)}</a></p>")
        else:
            h.append(f"<p>{escape(lab)}: {escape(v)}</p>")
    pattern = rng.randint(0, 4)
    personal = personal_items(p, include_contact=False)
    blocks = ([("personal", [("kv", lab, v) for _, lab, v in personal])] if personal else []) + \
        [(sec, section_lines(sec, p, rng, pattern)) for sec in section_order(p, rng)]
    for sec, lines in blocks:
        h.append(f"<h2>{escape(p.heading[sec])}</h2>")
        in_list = False
        for ln in lines:
            if ln[0] == "bullet":
                if not in_list:
                    h.append("<ul>")
                    in_list = True
                h.append(f"<li>{escape(ln[1])}</li>")
                continue
            if in_list:
                h.append("</ul>")
                in_list = False
            if ln[0] == "kv":
                h.append(f"<p><b>{escape(ln[1])}:</b> {escape(ln[2])}</p>")
            elif ln[0] == "lr":
                h.append(f"<table><tr><td>{_plain(ln[1])}</td><td>{_plain(ln[2])}</td></tr></table>")
            elif ln[0] == "title":
                h.append(f"<h3>{escape(ln[1])}</h3>")
            elif ln[0] != "space":
                h.append(f"<p>{escape(ln[1])}</p>")
        if in_list:
            h.append("</ul>")
    h.append("</body></html>")
    Path(path).write_text("\n".join(h), encoding="utf-8")


# template -> (renderer, extension, weight for free-form, weight for form-like)
TEMPLATES = {
    "pdf_classic":  (render_pdf_classic, ".pdf", 30, 5),
    "pdf_sidebar":  (render_pdf_sidebar, ".pdf", 20, 0),
    "pdf_form":     (render_pdf_form, ".pdf", 2, 40),
    "docx_classic": (render_docx_classic, ".docx", 20, 5),
    "docx_layout":  (render_docx_layout, ".docx", 10, 0),
    "docx_form":    (render_docx_form, ".docx", 2, 40),
    "txt":          (render_txt, ".txt", 5, 3),
    "md":           (render_md, ".md", 3, 2),
    "html":         (render_html, ".html", 3, 2),
}


def pick_template(p: Profile, rng: random.Random, only: list[str] | None = None) -> str:
    names = [n for n in TEMPLATES if not only or n in only]
    col = 3 if p.style.get("form_like") else 2
    weights = [TEMPLATES[n][col] or 0.5 for n in names]
    return rng.choices(names, weights=weights)[0]


def render(p: Profile, out_dir: str, stem: str, rng: random.Random, template: str | None = None) -> tuple[str, str]:
    template = template or pick_template(p, rng)
    fn, ext, _, _ = TEMPLATES[template]
    path = str(Path(out_dir) / f"{stem}{ext}")
    fn(p, path, rng)
    return path, template
