"""
Loads field definitions and picklist values from Excel files into structured data
for use by the CV parser pipeline.
"""
import openpyxl
from dataclasses import dataclass, field
from pathlib import Path

BASE_DIR = Path(__file__).parent


@dataclass
class FieldDef:
    field_id: str
    field_type: str  # String, Picklist, DateTime, Boolean
    label: str
    picklist_id: str | None
    required: bool
    remarks: str
    sample_value: str
    section: str | None  # e.g. "education", "awards", None for top-level
    nested_key: str | None  # e.g. "degree", "name" — the key inside results


@dataclass
class PicklistValue:
    external_code: str
    label_vi: str
    label_en: str
    option_id: str
    status: str


@dataclass
class Schema:
    fields: list[FieldDef] = field(default_factory=list)
    picklists: dict[str, list[PicklistValue]] = field(default_factory=dict)
    top_level_fields: list[FieldDef] = field(default_factory=list)
    sections: dict[str, list[FieldDef]] = field(default_factory=dict)

    def get_picklist_by_label(self, picklist_id: str, label: str) -> PicklistValue | None:
        values = self.picklists.get(picklist_id, [])
        label_lower = label.lower().strip()
        for v in values:
            if v.label_en.lower() == label_lower or v.label_vi.lower() == label_lower:
                return v
        return None

    def get_picklist_labels(self, picklist_id: str) -> list[str]:
        values = self.picklists.get(picklist_id, [])
        labels = []
        for v in values:
            if v.status == "A":
                labels.append(v.label_en)
                if v.label_vi != v.label_en:
                    labels.append(v.label_vi)
        return labels


def load_field_definitions(filepath: Path | None = None) -> list[FieldDef]:
    filepath = filepath or BASE_DIR / "Candidate Profile Fields JSON_260921.xlsx"
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    fields = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        field_id = row[0]
        if not field_id:
            continue

        parts = str(field_id).split(".")
        if len(parts) >= 3:
            section = parts[0]
            nested_key = parts[2]
        else:
            section = None
            nested_key = None

        fields.append(FieldDef(
            field_id=str(field_id),
            field_type=str(row[1] or "String"),
            label=str(row[2] or ""),
            picklist_id=str(row[3]) if row[3] else None,
            required=str(row[4]).upper() == "Y",
            remarks=str(row[5] or ""),
            sample_value=str(row[6] or ""),
            section=section,
            nested_key=nested_key,
        ))

    wb.close()
    return fields


def load_picklists(filepath: Path | None = None) -> dict[str, list[PicklistValue]]:
    filepath = filepath or BASE_DIR / "Candidate Profile Picklists_v2_260921.xlsx"
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    picklists: dict[str, list[PicklistValue]] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        pid = row[1]
        if not pid:
            continue
        pid = str(pid)
        if pid not in picklists:
            picklists[pid] = []

        picklists[pid].append(PicklistValue(
            external_code=str(row[3] or ""),
            label_vi=str(row[5] or ""),
            label_en=str(row[6] or ""),
            option_id=str(int(row[8])) if row[8] else "",
            status=str(row[7] or "A"),
        ))

    wb.close()
    return picklists


def load_schema() -> Schema:
    fields = load_field_definitions()
    picklists = load_picklists()

    top_level = []
    sections: dict[str, list[FieldDef]] = {}
    for f in fields:
        if f.section:
            sections.setdefault(f.section, []).append(f)
        else:
            top_level.append(f)

    return Schema(
        fields=fields,
        picklists=picklists,
        top_level_fields=top_level,
        sections=sections,
    )


if __name__ == "__main__":
    schema = load_schema()
    print(f"Loaded {len(schema.fields)} fields")
    print(f"  Top-level: {len(schema.top_level_fields)}")
    print(f"  Sections: {list(schema.sections.keys())}")
    print(f"Loaded {len(schema.picklists)} picklists:")
    for pid, vals in schema.picklists.items():
        active = [v for v in vals if v.status == "A"]
        print(f"  {pid}: {len(active)} active values")
