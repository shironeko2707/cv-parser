import unicodedata
import zipfile

import docx
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from text_extractor import extract
from text_utils import clean_line
from training.render import pdf_fonts

W, H = A4


@pytest.fixture(scope="module")
def font():
    return pdf_fonts()[0]


def _draw(path, font, items, pages=1, footer=None):
    """items: (page, x, y_from_top, text, size, bold)"""
    c = canvas.Canvas(str(path), pagesize=A4)
    for page in range(pages):
        for p, x, y, text, size, bold in items:
            if p == page:
                c.setFont(f"{font}-Bold" if bold else font, size)
                c.drawString(x, H - y, text)
        if footer:
            c.setFont(font, 8)
            c.drawString(40, 20, footer.format(n=page + 1))
        c.showPage()
    c.save()


def test_two_column_reading_order(tmp_path, font):
    items = [(0, 40, 40, "NGUYỄN VĂN AN", 20, True)]
    left = ["LIÊN HỆ", "an.nguyen@gmail.com", "0912 345 678", "Hà Nội", "KỸ NĂNG", "Python", "SQL", "Docker"]
    right = ["KINH NGHIỆM LÀM VIỆC", "Công ty FPT Software", "Lập trình viên", "01/2020 - nay",
             "HỌC VẤN", "Đại học Bách khoa Hà Nội", "2015 - 2019", "Kỹ sư CNTT"]
    for i, (lt, rt) in enumerate(zip(left, right)):  # same baselines on purpose
        y = 90 + i * 14
        items.append((0, 40, y, lt, 10, lt.isupper()))
        items.append((0, 250, y, rt, 10, rt.isupper()))
    path = tmp_path / "two_col.pdf"
    _draw(path, font, items)
    text = extract(path).layout_text
    assert text.index("Docker") < text.index("Công ty FPT Software"), text  # left column first
    assert "## KINH NGHIỆM LÀM VIỆC" in text and "## HỌC VẤN" in text
    assert "0912 345 678 |" not in text  # columns not merged into rows


def test_right_aligned_dates_stay_on_row(tmp_path, font):
    items = [(0, 40, 40, "John Smith", 20, True), (0, 40, 80, "EXPERIENCE", 12, True)]
    for i, (title, dates) in enumerate([("Senior Engineer, Acme", "Jan 2020 - Present"),
                                        ("Engineer, Globex", "Mar 2017 - Dec 2019"),
                                        ("Intern, Initech", "Jun 2016 - Feb 2017"),
                                        ("Tutor, Uni", "2014 - 2016"), ("Barista, Cafe", "2013 - 2014")]):
        y = 100 + i * 40
        items += [(0, 40, y, title, 10, True), (0, 440, y, dates, 10, False),
                  (0, 40, y + 13, f"Did useful things number {i}", 10, False)]
    path = tmp_path / "dates.pdf"
    _draw(path, font, items)
    text = extract(path).layout_text
    assert "Senior Engineer, Acme | Jan 2020 - Present" in text, text
    assert "Engineer, Globex | Mar 2017 - Dec 2019" in text


def test_dates_are_not_page_numbers(tmp_path, font):
    items = [(0, 40, 40, "Jane Doe", 18, True), (0, 40, 80, "EDUCATION", 12, True),
             (0, 40, 100, "Hanoi University", 10, False), (0, 40, 114, "09/2012", 10, False),
             (0, 40, 128, "3", 10, False)]
    path = tmp_path / "noise.pdf"
    _draw(path, font, items)
    text = extract(path).layout_text
    assert "09/2012" in text
    assert "\n3\n" not in f"\n{text}\n"  # bare page number removed


def test_repeated_header_footer_removed(tmp_path, font):
    items = []
    for page in range(2):
        items += [(page, 40, 20, "CV - Trần Thị Bình - Confidential", 8, False),
                  (page, 40, 100, f"Body text on page {page + 1}", 10, False)]
    path = tmp_path / "hf.pdf"
    _draw(path, font, items, pages=2, footer="Page {n} of 2")
    text = extract(path).layout_text
    assert text.count("Confidential") == 1
    assert "Page 1 of 2" not in text and "Body text on page 2" in text


def test_pdf_ruled_table_rows(tmp_path):
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    f = pdf_fonts()[0]
    path = tmp_path / "table.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    t = Table([["Từ", "Đến", "Công ty", "Chức vụ"],
               ["01/2020", "nay", "Viettel", "Kỹ sư"],
               ["06/2017", "12/2019", "FPT", ""]])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "black"), ("FONTNAME", (0, 0), (-1, -1), f)]))
    doc.build([t])
    d = extract(path)
    assert "01/2020 | nay | Viettel | Kỹ sư" in d.layout_text
    assert d.layout_text.count("Viettel") == 1  # not duplicated as loose text
    assert d.tables and d.tables[0][0][0] == "Từ"


def test_docx_header_layout_table_bullets_and_textbox(tmp_path):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Lê Minh Châu | chau.le@gmail.com"
    d.add_paragraph("KINH NGHIỆM LÀM VIỆC")
    layout = d.add_table(rows=1, cols=2)
    left, right = layout.rows[0].cells
    left.text = "KỸ NĂNG"
    for s in ("Excel", "SAP", "Power BI"):
        left.add_paragraph(s)
    right.text = "Công ty Vinamilk"
    for s in ("Kế toán viên", "03/2019 - nay", "Lập báo cáo tài chính"):
        right.add_paragraph(s)
    data = d.add_table(rows=2, cols=3)
    for j, v in enumerate(["Họ tên", "Năm sinh", "Nghề nghiệp"]):
        data.rows[0].cells[j].text = v
    data.rows[1].cells[0].text = "Lê Văn Hùng"
    data.rows[1].cells[2].text = "Giáo viên"
    d.add_paragraph("Tiếng Anh - IELTS 7.0", style="List Bullet")
    # text box (common in CV templates; python-docx paragraph.text misses it)
    tb_xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<w:pict><wps:txbx><w:txbxContent><w:p><w:r><w:t>Textbox: 0987 654 321</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></w:pict></w:r>'
    )
    from lxml import etree
    d.add_paragraph()._p.append(etree.fromstring(tb_xml))
    path = tmp_path / "cv.docx"
    d.save(path)

    text = extract(path).layout_text
    assert text.startswith("Lê Minh Châu | chau.le@gmail.com")
    assert "## KINH NGHIỆM LÀM VIỆC" in text
    assert text.index("Power BI") < text.index("Công ty Vinamilk")
    assert "Lê Văn Hùng |  | Giáo viên" in text  # empty cell keeps column alignment
    assert "• Tiếng Anh - IELTS 7.0" in text
    assert "Textbox: 0987 654 321" in text


def test_plain_formats(tmp_path):
    (tmp_path / "cv.txt").write_text("PHẠM THU HÀ\nha.pham@yahoo.com\n\nHỌC VẤN\n- Đại học Ngoại thương, 2018\n",
                                     encoding="utf-8")
    t = extract(tmp_path / "cv.txt").layout_text
    assert "## HỌC VẤN" in t and "• Đại học Ngoại thương, 2018" in t

    (tmp_path / "cv.html").write_text("<h1>Jane Doe</h1><h2>Education</h2><ul><li>MIT, 2010</li></ul>"
                                      "<a href='https://linkedin.com/in/jane'>in</a>", encoding="utf-8")
    h = extract(tmp_path / "cv.html")
    assert "## Education" in h.layout_text and h.links == ["https://linkedin.com/in/jane"]

    rtf = (b"{\\rtf1\\ansi\\ansicpg1252{\\fonttbl{\\f0 Arial;}}{\\b Tr\\u7847?n V\\u259?n B}"
           b"\\par Email: b@x.vn\\par}")
    (tmp_path / "cv.rtf").write_bytes(rtf)
    r = extract(tmp_path / "cv.rtf").layout_text
    assert "Trần Văn B" in r and "b@x.vn" in r

    with zipfile.ZipFile(tmp_path / "cv.odt", "w") as z:
        z.writestr("content.xml", (
            '<?xml version="1.0"?><office:document-content '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            '<office:body><office:text><text:h>Kinh nghiệm</text:h><text:p>Viettel<text:tab/>2020</text:p>'
            '</office:text></office:body></office:document-content>'))
    o = extract(tmp_path / "cv.odt").layout_text
    assert "## Kinh nghiệm" in o and "Viettel | 2020" in o


def test_clean_line_unicode():
    decomposed = unicodedata.normalize("NFD", "Nguyễn Thị Hương")
    assert clean_line(decomposed) == "Nguyễn Thị Hương"
    assert clean_line(" an@x.vn\x00") == "an@x.vn"  # icon font + unmapped glyph
    assert clean_line("▪ Quản lý dự án") == "• Quản lý dự án"
    assert clean_line("ofﬁce") == "office"
